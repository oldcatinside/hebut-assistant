"""表头别名、保守自动识别及用户指定的逐 Sheet 映射。"""
from dataclasses import dataclass, field
import re
import unicodedata
from openpyxl.utils import get_column_letter
from app.phone_utils import text_value

FIELDS = ('姓名', '公司名称', '宣讲会时间', '宣讲会地点', '联系人', '联系人电话')
EXTRA_FIELDS = ('结束时间', '宣讲日期')
ALIASES = {
    '姓名': ('姓名', '负责人', '接待人', '接待人员', '接待学生', '负责学生', '学生姓名', '接待部负责人', '接待姓名', '值班人员', '分配人员'),
    '公司名称': ('公司名称', '企业名称', '单位名称', '招聘单位', '用人单位', '宣讲企业', '企业'),
    '宣讲会时间': ('宣讲会时间', '宣讲时间', '开始时间', '宣讲开始时间', '宣讲会开始时间', '时间'),
    '宣讲会地点': ('宣讲会地点', '宣讲地点', '举办地点', '场地', '教室', '地点'),
    '联系人': ('联系人', '企业联系人', '单位联系人', '联系人姓名'),
    '联系人电话': ('联系人电话', '联系电话', '企业联系电话', '联系人手机号', '手机号码', '手机号', '电话'),
    '结束时间': ('结束时间', '宣讲会结束时间', '宣讲结束时间'),
    '宣讲日期': ('宣讲日期', '宣讲会日期', '日期'),
}


def header_key(value):
    return re.sub(r'[\s\u200b\ufeff:：*＊]', '', unicodedata.normalize('NFKC', text_value(value)))


LOOKUP = {header_key(alias): field for field, aliases in ALIASES.items() for alias in aliases}


@dataclass
class SheetLayout:
    mode: str = 'auto'
    columns: dict[str, int] = field(default_factory=dict)
    start_row: int = 1
    end_row: int = 0
    merged_names: bool = False
    year: int = 0


def header_mapping(sheet, row):
    mapping, duplicates = {}, []
    for col in range(sheet.width):
        current = header_key(sheet.value(row, col, merged=True))
        above = header_key(sheet.value(row-1, col, merged=True)) if row else ''
        key = LOOKUP.get(above + current) if above and current and above != current else None
        key = key or LOOKUP.get(current)
        if key:
            if key in mapping:
                duplicates.append(key)
            else:
                mapping[key] = col
    return mapping, duplicates


def detect_sections(sheet, name):
    headers = []
    for row in range(len(sheet.rows)):
        mapping, duplicates = header_mapping(sheet, row)
        if ('姓名' in mapping and (len(mapping) >= 2 or len(sheet.rows[row]) == 1)) or len(mapping) >= 2:
            if headers and row == headers[-1][0] + 1:
                combined = dict(headers[-1][1])
                for key, col in mapping.items():
                    combined = {k: v for k, v in combined.items() if v != col or k == key}
                    combined[key] = col
                headers[-1] = (row, combined, duplicates)
            else:
                headers.append((row, mapping, duplicates))
    sections, warnings = [], []
    for position, (header, mapping, duplicates) in enumerate(headers):
        end = headers[position+1][0] if position+1 < len(headers) else len(sheet.rows)
        if duplicates:
            warnings.append(f'第{header+1}行有重复字段：{"、".join(duplicates)}；该段未读取，请手动指定列。')
            continue
        mapping = dict(mapping)
        if '姓名' not in mapping:
            candidates = [col for col in range(sheet.width)
                          if not header_key(sheet.value(header, col, merged=True))
                          and col not in mapping.values()
                          and any(text_value(sheet.value(row, col)) == name for row in range(header+1, end))]
            if len(candidates) == 1 and sum(key in mapping for key in FIELDS[1:]) >= 3:
                mapping['姓名'] = candidates[0]
                warnings.append(f'第{header+1}行没有姓名表头，自动采用 {get_column_letter(candidates[0]+1)} 列（唯一空表头列精确匹配筛选姓名），请在“列设置 / 原表预览”核对。')
            else:
                details = '、'.join(get_column_letter(c+1) for c in candidates) or '无可靠候选'
                warnings.append(f'第{header+1}行缺少明确的姓名列（{details}）；该段未读取，请手动指定列。')
                continue
        sections.append(SheetLayout('manual', mapping, header+2, end))
    if not headers:
        warnings.append('没有可用表头；请打开“列设置 / 原表预览”，指定姓名列、其他信息列与数据起始行。')
    return sections, warnings
