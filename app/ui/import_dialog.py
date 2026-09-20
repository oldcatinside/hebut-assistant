from copy import deepcopy
from PySide6.QtCore import Qt
from PySide6.QtGui import QColor
from PySide6.QtWidgets import (QDialog, QVBoxLayout, QHBoxLayout, QLabel, QComboBox, QSpinBox,
    QCheckBox, QGridLayout, QTableWidget, QTableWidgetItem, QDialogButtonBox, QMessageBox, QAbstractItemView)
from openpyxl.utils import get_column_letter
from app.import_layout import FIELDS, EXTRA_FIELDS, SheetLayout
from app.phone_utils import text_value


class ImportDialog(QDialog):
    def __init__(self, book, name, layouts=None, suggested=None, parent=None):
        super().__init__(parent)
        self.setWindowTitle('列设置 / 原表预览')
        self.resize(1100, 760)
        self.book, self.name = book, name
        self.layouts = deepcopy(layouts or {})
        self.suggested = suggested or {}
        self.active = None
        root = QVBoxLayout(self)
        info = QLabel('每个工作表分别设置。完全没有表头也可指定列；姓名仍然精确匹配。设置只对当前文件生效，不修改原表。')
        info.setWordWrap(True)
        root.addWidget(info)
        bar = QHBoxLayout()
        bar.addWidget(QLabel('工作表：'))
        self.sheet_choice = QComboBox()
        self.sheet_choice.addItems([s.name for s in book.sheets])
        bar.addWidget(self.sheet_choice, 1)
        bar.addWidget(QLabel('读取方式：'))
        self.mode = QComboBox()
        self.mode.addItem('自动识别', 'auto')
        self.mode.addItem('手动指定列和行', 'manual')
        self.mode.addItem('跳过本工作表', 'skip')
        bar.addWidget(self.mode)
        root.addLayout(bar)
        grid = QGridLayout()
        self.columns = {}
        for i, field in enumerate((*FIELDS, *EXTRA_FIELDS)):
            combo = QComboBox()
            combo.setMinimumWidth(180)
            self.columns[field] = combo
            grid.addWidget(QLabel(field + '：'), i//2, (i%2)*2)
            grid.addWidget(combo, i//2, (i%2)*2+1)
            combo.currentIndexChanged.connect(self.update_match_count)
        root.addLayout(grid)
        ranges = QHBoxLayout()
        self.start, self.end, self.year = QSpinBox(), QSpinBox(), QSpinBox()
        self.year.setRange(0, 9999)
        self.year.setSpecialValueText('不补年份')
        for title, widget in (('数据起始行：', self.start), ('结束行：', self.end), ('缺失年份时使用：', self.year)):
            ranges.addWidget(QLabel(title))
            ranges.addWidget(widget)
        root.addLayout(ranges)
        self.merged = QCheckBox('展开 Excel 中实际合并的姓名单元格（普通空白不向下填充）')
        root.addWidget(self.merged)
        hint = QLabel('“宣讲日期”仅用于日期与开始钟点分列；完整日期时间请只指定“宣讲会时间”。缺少年份时必须由您明确填写年份。')
        hint.setWordWrap(True)
        root.addWidget(hint)
        self.match_count = QLabel()
        self.match_count.setTextFormat(Qt.PlainText)
        root.addWidget(self.match_count)
        preview_bar = QHBoxLayout()
        preview_bar.addWidget(QLabel('原始单元格预览（每页最多200行，黄色为精确匹配姓名）：从第'))
        self.preview_start = QSpinBox()
        preview_bar.addWidget(self.preview_start)
        preview_bar.addWidget(QLabel('行开始'))
        preview_bar.addStretch()
        root.addLayout(preview_bar)
        self.preview = QTableWidget()
        self.preview.setEditTriggers(QAbstractItemView.NoEditTriggers)
        root.addWidget(self.preview, 1)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText('应用并重新读取')
        buttons.button(QDialogButtonBox.Cancel).setText('取消')
        buttons.accepted.connect(self.validate_accept)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)
        self.sheet_choice.currentIndexChanged.connect(self.change_sheet)
        self.preview_start.valueChanged.connect(self.render_preview)
        self.mode.currentIndexChanged.connect(self.update_mode)
        self.start.valueChanged.connect(self.update_match_count)
        self.end.valueChanged.connect(self.update_match_count)
        self.merged.toggled.connect(self.update_match_count)
        self.change_sheet(0)

    def save_active(self):
        if self.active is not None:
            sheet = self.book.sheets[self.active]
            self.layouts[sheet.name] = SheetLayout(self.mode.currentData(),
                {field: combo.currentData() for field, combo in self.columns.items() if combo.currentData() is not None},
                self.start.value(), self.end.value(), self.merged.isChecked(), self.year.value())

    def change_sheet(self, index):
        self.save_active()
        self.active = None
        sheet = self.book.sheets[index]
        setting = self.layouts.get(sheet.name)
        if setting is None:
            setting = deepcopy(self.suggested.get(sheet.name, SheetLayout()))
            setting.mode = 'auto'
        self.mode.setCurrentIndex(self.mode.findData(setting.mode))
        for field, combo in self.columns.items():
            combo.clear()
            combo.addItem('未指定 / 此字段缺失', None)
            for col in range(sheet.width):
                samples = [text_value(sheet.value(r, col)) for r in range(min(len(sheet.rows), 20)) if text_value(sheet.value(r, col))]
                label = ' | '.join(samples[:2])[:42]
                combo.addItem(f'{get_column_letter(col+1)} 列 · {label}', col)
            combo.setCurrentIndex(max(0, combo.findData(setting.columns.get(field))))
        for spin in (self.start, self.end, self.preview_start):
            spin.setRange(1, max(1, len(sheet.rows)))
        self.start.setValue(setting.start_row)
        self.end.setValue(setting.end_row or max(1, len(sheet.rows)))
        self.year.setValue(setting.year)
        self.merged.setChecked(setting.merged_names)
        self.preview_start.setValue(1)
        self.active = index
        self.update_mode()
        self.render_preview()

    def update_mode(self):
        enabled = self.mode.currentData() == 'manual'
        for widget in (*self.columns.values(), self.start, self.end, self.year, self.merged):
            widget.setEnabled(enabled)
        self.update_match_count()

    def update_match_count(self):
        if self.active is None:
            return
        sheet = self.book.sheets[self.active]
        col = self.columns['姓名'].currentData()
        if col is None:
            self.match_count.setText('未指定姓名列；自动模式将尝试识别，手动模式必须指定。')
            return
        count = sum(text_value(sheet.value(r, col, self.merged.isChecked())) == self.name
                    for r in range(self.start.value()-1, self.end.value()))
        self.match_count.setText(f'当前列设置：{get_column_letter(col+1)} 列，姓名“{self.name}”精确匹配 {count} 行。请核对这确实是接待姓名，而非企业联系人。')

    def render_preview(self):
        if self.active is None:
            return
        sheet = self.book.sheets[self.active]
        first = self.preview_start.value()-1
        rows = sheet.rows[first:first+200]
        self.preview.setRowCount(len(rows))
        self.preview.setColumnCount(sheet.width)
        self.preview.setHorizontalHeaderLabels([get_column_letter(c+1) for c in range(sheet.width)])
        self.preview.setVerticalHeaderLabels([str(first+r+1) for r in range(len(rows))])
        for ri, row in enumerate(rows):
            for ci in range(sheet.width):
                value = text_value(row[ci]) if ci < len(row) else ''
                item = QTableWidgetItem(value)
                item.setToolTip(value)
                if value == self.name:
                    item.setBackground(QColor('#fff0a6'))
                self.preview.setItem(ri, ci, item)
        for col in range(sheet.width):
            self.preview.setColumnWidth(col, 180)

    def validate_accept(self):
        self.save_active()
        for sheet in self.book.sheets:
            setting = self.layouts.get(sheet.name, SheetLayout())
            if setting.mode != 'manual':
                continue
            if '姓名' not in setting.columns or setting.end_row < setting.start_row:
                QMessageBox.warning(self, '设置不完整', f'工作表「{sheet.name}」必须指定姓名列及有效行范围。')
                return
            if len(set(setting.columns.values())) != len(setting.columns):
                QMessageBox.warning(self, '列重复', f'工作表「{sheet.name}」不能把同一列指定给多个字段。')
                return
        self.accept()
