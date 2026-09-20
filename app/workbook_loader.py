"""按内容选择读取引擎，只读原文件并保留实际合并区域。"""
from dataclasses import dataclass, field
from functools import cached_property
from datetime import datetime
from pathlib import Path
from zipfile import is_zipfile, ZipFile
import xlrd
from openpyxl import load_workbook
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH

MAX_CELLS = 2_000_000


@dataclass
class RawSheet:
    name: str
    rows: list[list]
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)
    hidden: bool = False

    @cached_property
    def width(self):
        return max((len(r) for r in self.rows), default=0)

    def value(self, row, col, merged=False):
        if merged:
            for r0, r1, c0, c1 in self.merges:
                if r0 <= row < r1 and c0 <= col < c1:
                    row, col = r0, c0
                    break
        return self.rows[row][col] if 0 <= row < len(self.rows) and 0 <= col < len(self.rows[row]) else None


@dataclass
class RawBook:
    sheets: list[RawSheet]
    epoch: datetime
    warnings: list[str] = field(default_factory=list)


def load_raw_book(path) -> RawBook:
    path = Path(path)
    if not path.is_file():
        raise ValueError('文件不存在，请重新选择 Excel 文件')
    with path.open('rb') as handle:
        signature = handle.read(512)
    warnings = []
    try:
        if not signature.startswith(b'\xd0\xcf\x11\xe0') and is_zipfile(path):
            with ZipFile(path) as archive:
                if 'xl/workbook.xml' not in archive.namelist():
                    raise ValueError('该文件不是受支持的 Excel 工作簿（可能是 xlsb 或普通压缩包）。请用 Excel/WPS 另存为 .xlsx。')
            if path.suffix.lower() not in ('.xlsx', '.xlsm'):
                warnings.append('扩展名与内容不符，已按实际的新版 Excel 格式读取。')
            with path.open('rb') as stream:
                book = load_workbook(stream, data_only=True, read_only=False)
                sheets, total = [], 0
                try:
                    for sheet in book.worksheets:
                        total += sheet.max_row * sheet.max_column
                        if total > MAX_CELLS:
                            raise ValueError('工作簿使用区域过大（超过200万格）。请删除末尾空白格式或拆分工作簿后重试。')
                        sheets.append(RawSheet(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)],
                            [(m.min_row-1, m.max_row, m.min_col-1, m.max_col) for m in sheet.merged_cells.ranges], sheet.sheet_state != 'visible'))
                    return RawBook(sheets, book.epoch, warnings)
                finally:
                    book.close()
        if signature.startswith(b'\xd0\xcf\x11\xe0'):
            if path.suffix.lower() != '.xls':
                warnings.append('扩展名与内容不符，已按实际的旧版 Excel 格式读取。')
            book = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
            try:
                sheets, total = [], 0
                for sheet in book.sheets():
                    total += sheet.nrows * sheet.ncols
                    if total > MAX_CELLS:
                        raise ValueError('工作簿使用区域过大（超过200万格）。请删除末尾空白格式或拆分工作簿后重试。')
                    rows = []
                    for ri in range(sheet.nrows):
                        row = []
                        for cell in sheet.row(ri):
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                            elif cell.ctype == xlrd.XL_CELL_ERROR:
                                row.append('#Excel单元格错误')
                            else:
                                row.append(cell.value)
                        rows.append(row)
                    sheets.append(RawSheet(sheet.name, rows, sheet.merged_cells, bool(sheet.visibility)))
                return RawBook(sheets, MAC_EPOCH if book.datemode else WINDOWS_EPOCH, warnings)
            finally:
                book.release_resources()
        if not signature.strip():
            raise ValueError('文件为空，请确认下载或保存已完成。')
        if b'<' in signature or b'\x00<' in signature:
            raise ValueError('文件实际可能是网页表格或 XML，并非真正的 Excel。请用 Excel/WPS 打开，再“另存为 Excel 工作簿（.xlsx）”，不要只改扩展名。')
        raise ValueError('无法识别文件内容。若是 CSV/文本，请用 Excel/WPS 导入后另存为 .xlsx；若文件损坏，请重新下载或使用“打开并修复”。')
    except (PermissionError, ValueError):
        raise
    except Exception as exc:
        raise ValueError('工作簿无法解析，可能已加密、损坏或格式不兼容。请用 Excel/WPS 打开：加密文件需先解密；损坏文件尝试“打开并修复”；再另存为 .xlsx 后重试。') from exc
