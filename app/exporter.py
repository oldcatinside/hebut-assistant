import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from app.models import Record


def export_records(records: list[Record], target, source=None) -> Path:
    target = Path(target).resolve()
    if target.suffix.lower() != '.xlsx':
        raise ValueError('导出文件必须使用 .xlsx 扩展名')
    if source:
        source = Path(source).resolve()
        if target == source or (target.exists() and source.exists() and os.path.samefile(target, source)):
            raise ValueError('不允许覆盖原始 Excel 文件，请选择其他文件名')
    wb = Workbook()
    ws = wb.active
    ws.title = '宣讲短信整理'
    headers = ['公司名称', '联系人', '联系电话', '格式化宣讲时间', '宣讲地点', '生成短信', '数据状态', '短信生成状态', '来源工作表', '来源行', '原始宣讲时间', '原始地点', '原始联系电话']
    ws.append(headers)
    for record in records:
        values = [record.company, record.contact, record.phone, record.time_text, record.address, record.message,
                  record.status, record.message_error or '已生成', record.sheet, str(record.row), record.raw_time, record.raw_address, record.raw_phone]
        ws.append(values)
        for cell in ws[ws.max_row]:
            # 所有导出内容为纯文本，企业名称以等号开头时也不执行公式。
            cell.data_type = 's'
            cell.number_format = '@'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        if record.issues:
            ws.cell(ws.max_row, 7).font = Font(color='B42318')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='245A81')
    from openpyxl.utils import get_column_letter
    for index, width in enumerate([34, 16, 20, 34, 34, 90, 42, 32, 22, 12, 34, 34, 22], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    temp_path = None
    try:
        with NamedTemporaryFile(dir=target.parent, suffix='.xlsx', delete=False) as handle:
            temp_path = Path(handle.name)
        wb.save(temp_path)
        os.replace(temp_path, target)
    finally:
        wb.close()
        if temp_path and temp_path.exists():
            temp_path.unlink()
    return target
